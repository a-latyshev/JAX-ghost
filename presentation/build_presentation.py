"""Editable PowerPoint and vector PDF from the same positioned slide objects.

Run with the isolated presentation environment; no FEM/GPU runtime is imported.
Coordinates and typography use a 1280 x 720 design canvas, scaled to 13⅓ x 7½ in.
"""
from pathlib import Path
import json
import math
import os

import pymupdf as fitz
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor

ROOT = Path(__file__).resolve().parent
W, H, SCALE = 1280, 720, 0.75
SLIDE_COUNT = 8
INK, MUTED, LINE = '#172638', '#546477', '#D9E1EA'
BLUE, ORANGE, TEAL = '#2463C9', '#C8531C', '#087D70'
PALE_B, PALE_O, PALE_T = '#EFF5FF', '#FFF3EB', '#EDF8F5'
BG, WHITE = '#F5F7FA', '#FFFFFF'
FONTS = Path(os.environ.get('PRESENTATION_FONT_DIR', str(ROOT / 'fonts')))
for name, filename in [('regular', 'DejaVuSans.ttf'), ('bold', 'DejaVuSans-Bold.ttf'),
                       ('mono', 'DejaVuSansMono.ttf')]:
    pdfmetrics.registerFont(TTFont(name, str(FONTS / filename)))


class Slide:
    def __init__(self, number, title, topic, subtitle):
        self.number, self.title, self.objects = number, title, []
        self.notes = ''
        self.text(56, 26, 850, 'JAX-GHOST   /   ' + topic.upper(), 14, MUTED, bold=True)
        self.rect(1173, 26, 51, 5, BLUE)
        self.text(56, 63, 1170, title, 35, INK, bold=True)
        self.text(56, 120, 1170, subtitle, 18, MUTED)
        self.line(56, 674, 1224, 674, LINE, 1)
        self.text(56, 688, 720, 'FEniCSx → MPI ownership → JAX GPU communication', 12, MUTED)
        for i in range(SLIDE_COUNT):
            self.rect(1030 + i * 19, 694, 12, 4, TEAL if i == number else LINE)
        self.text(1200, 683, 24, str(number), 19, INK, bold=True, align='right')

    def add(self, kind, **kw):
        self.objects.append(dict(kind=kind, **kw))

    def text(self, x, y, w, text, size=22, color=INK, bold=False, font=None, align='left'):
        font = font or ('bold' if bold else 'regular')
        lines = text.split('\n')
        for t in lines:
            width = pdfmetrics.stringWidth(t, font, size)
            if width > w - 2:
                raise ValueError(f'Slide {self.number}, text too wide ({width:.1f}>{w}): {t}')
        self.add('text', x=x, y=y, w=w, h=size * 1.3 * len(lines) + 4,
                 text=text, size=size, color=color, font=font, align=align)

    def rect(self, x, y, w, h, fill=WHITE, stroke=None, lw=1.5, radius=0):
        self.add('rect', x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, lw=lw, radius=radius)

    def circle(self, x, y, r, fill, stroke=None, lw=2):
        self.add('ellipse', x=x-r, y=y-r, w=2*r, h=2*r, fill=fill, stroke=stroke, lw=lw)

    def poly(self, points, fill, stroke=None, lw=1):
        self.add('poly', points=points, fill=fill, stroke=stroke, lw=lw)

    def line(self, x1, y1, x2, y2, color=LINE, lw=2, dash=False):
        self.add('line', x1=x1, y1=y1, x2=x2, y2=y2, color=color, lw=lw, dash=dash)

    def arrow(self, x1, y1, x2, y2, color=MUTED, lw=2.5, head=10):
        self.line(x1, y1, x2, y2, color, lw)
        a = math.atan2(y2-y1, x2-x1)
        self.poly([(x2, y2), (x2-head*math.cos(a)+head*.48*math.sin(a), y2-head*math.sin(a)-head*.48*math.cos(a)),
                   (x2-head*math.cos(a)-head*.48*math.sin(a), y2-head*math.sin(a)+head*.48*math.cos(a))], color)

    def takeaway(self, text, small=None):
        self.rect(56, 604, 1168, 57, PALE_T, radius=8)
        self.rect(56, 604, 5, 57, TEAL)
        self.text(76, 617 if small is None else 609, 1120, text, 21, TEAL, bold=True)
        if small:
            self.text(76, 637, 1120, small, 14, MUTED)


def chip(s, x, y, w, label, color, fill=None, size=17):
    s.rect(x, y, w, 34, fill or WHITE, color, 1.2, radius=6)
    s.text(x+10, y+5, w-20, label, size, color, bold=True, align='center')


def box(s, x, y, w, h, title, body=None, color=INK, fill=BG, size=21):
    s.rect(x, y, w, h, fill, LINE, 1, radius=10)
    s.text(x+20, y+17, w-40, title, size, color, bold=True)
    if body:
        s.text(x+20, y+54, w-40, body, 18, MUTED)


def vector(s, x, y, labels, owners, ghosts=(), padding=(), cell=40, h=43, fs=19):
    for i, label in enumerate(labels):
        col = [BLUE, ORANGE][owners[i]] if owners[i] is not None else LINE
        pad, ghost = i in padding, i in ghosts
        s.rect(x+i*cell, y, cell-3, h, BG if pad else WHITE if ghost else col,
               LINE if pad else col, 2 if ghost else 1, radius=3)
        s.text(x+i*cell+2, y+(h-fs*1.3)/2-1, cell-7, label, fs,
               MUTED if pad else col if ghost else WHITE, bold=True, align='center')


OWNERS = {i: (1 if i % 4 == 3 or i == 6 else 0) for i in range(12)}
R0 = [0, 1, 2, 4, 5, 8, 9, 10, 6]
R1 = [3, 6, 7, 11, 2, 10]
SUB = str.maketrans('0123456789', '₀₁₂₃₄₅₆₇₈₉')
def u(i):
    return 'u' + str(i).translate(SUB)


def mesh(s, x, y, dx, dy, columns, rank=None, labels=True, r=12):
    points = {row*4+col: (x+(col-columns[0])*dx, y+row*dy)
              for row in range(3) for col in columns}
    for row in range(2):
        for col in columns[:-1]:
            ids = (row*4+col, row*4+col+1, (row+1)*4+col, (row+1)*4+col+1)
            a,b,c,d = [points[i] for i in ids]
            fill = PALE_B if col < 2 else PALE_O
            s.poly([a,b,d], fill, LINE, 1.5)
            s.poly([a,c,d], fill, LINE, 1.5)
    for i, (xx,yy) in points.items():
        col = [BLUE, ORANGE][OWNERS[i]]
        ghost = rank is not None and OWNERS[i] != rank
        s.circle(xx, yy, r, WHITE if ghost else col, col, 2)
        if labels:
            s.text(xx-r+1, yy-9, 2*r-2, str(i), 12, col if ghost else WHITE, bold=True, align='center')
    return points


def slides():
    out = []
    s = Slide(0, 'Finite elements → distributed algebra', 'The starting point',
              'A brief reminder: local basis functions turn a PDE into a sparse system.')
    for x, title in [(56, '01   Mesh the domain'), (461, '02   Approximate the field'), (866, '03   Assemble the system')]:
        s.rect(x, 184, 358, 330, BG, radius=12)
        s.text(x+24, 206, 310, title, 20, INK, bold=True)
    # A small neutral triangular mesh, with a highlighted element.
    for row in range(3):
        for col in range(4):
            x,y = 99+col*64, 276+row*55
            s.poly([(x,y),(x+64,y),(x+64,y+55)], PALE_B, '#95ACCA', 1)
            s.poly([(x,y),(x,y+55),(x+64,y+55)], PALE_B, '#95ACCA', 1)
    s.poly([(163,331),(227,331),(227,386)], BLUE, BLUE)
    s.text(81, 462, 310, 'Small elements, local support', 18, MUTED)
    s.text(481, 304, 318, 'uₕ = ∑ᵢ uᵢ φᵢ', 39, BLUE, align='center')
    s.text(485, 383, 310, 'DoFs uᵢ weight basis functions φᵢ.', 17, MUTED, align='center')
    s.text(485, 462, 310, 'For scalar P1: one DoF per vertex', 16, MUTED)
    s.text(891, 295, 304, 'Au = b', 51, TEAL, bold=True, align='center')
    for i in range(5):
        for j in range(5):
            if abs(i-j) <= 1:
                s.rect(984+j*18, 373+i*15, 12, 10, TEAL)
    s.text(890, 462, 310, 'Sparse coupling between DoFs', 17, MUTED)
    s.arrow(423, 344, 451, 344, MUTED, 3)
    s.arrow(828, 344, 856, 344, MUTED, 3)
    s.text(56, 550, 1140, 'FEniCSx: the platform     •     DOLFINx: meshes, function spaces and distributed assembly', 21, INK)
    s.takeaway('Local finite-element support creates sparse algebra with distributed ownership.')
    s.notes = '''Timing: 0:40.

Finite elements split the domain into cells and express the approximate solution
as a weighted sum of local basis functions. Those weights are the degrees of
freedom. For the scalar continuous P1 example used next, each vertex has one
DoF; that vertex picture does not apply to all finite elements. A variational
form and assembly produce a sparse algebraic system because basis functions
have local support. Boundary conditions and the specific PDE are omitted here.

FEniCSx is the finite-element computing platform; DOLFINx is its computational
core and problem-solving interface, handling meshes, function spaces and
distributed assembly. We retain its ownership conventions when moving numerical
operations into JAX. Transition: how are those DoFs distributed across ranks?

Reference: https://docs.fenicsproject.org/dolfinx/main/python/
'''
    out.append(s)

    s = Slide(1, 'One mesh, two MPI ranks', 'Ownership',
              'Scalar P1 • colors identify the owner • outlines identify local ghost copies')
    s.text(65, 182, 280, 'GLOBAL CELL PARTITION', 15, MUTED, bold=True)
    mesh(s, 88, 270, 62, 75, [0,1,2,3], r=11)
    s.text(63, 454, 288, 'Illustrative ownership;\nactual IDs come from IndexMap.', 16, MUTED)
    chip(s, 64, 523, 117, 'Rank 0', BLUE, PALE_B)
    chip(s, 193, 523, 117, 'Rank 1', ORANGE, PALE_O)
    s.arrow(318, 351, 377, 351)
    s.rect(398, 182, 382, 390, BG, radius=12)
    s.rect(840, 182, 384, 390, BG, radius=12)
    s.text(420, 200, 340, 'RANK 0   /   8 owned + 1 ghost', 18, BLUE, bold=True)
    s.text(860, 200, 344, 'RANK 1   /   4 owned + 2 ghosts', 18, ORANGE, bold=True)
    p0 = mesh(s, 450, 270, 100, 75, [0,1,2], rank=0)
    p1 = mesh(s, 965, 270, 145, 75, [2,3], rank=1)
    for node in [2,6,10]:
        x0,y0 = p0[node]; x1,y1 = p1[node]
        color = [BLUE,ORANGE][OWNERS[node]]
        if OWNERS[node] == 0:
            s.arrow(x0+16,y0,x1-18,y1,color,2)
        else:
            s.arrow(x1-16,y1,x0+18,y0,color,2)
        s.rect(740, y0-28, 116, 24, WHITE)
        s.text(744, y0-26, 108, 'copy '+u(node), 15, color, bold=True, align='center')
    vector(s, 425, 469, [u(i) for i in R0], [OWNERS[i] for i in R0], ghosts=[8], cell=37, fs=16)
    vector(s, 870, 469, [u(i) for i in R1], [OWNERS[i] for i in R1], ghosts=[4,5], cell=52, fs=18)
    s.text(425, 526, 320, '[ owned values | ghost copies ]', 17, MUTED)
    s.text(870, 526, 325, '[ owned values | ghost copies ]', 17, MUTED)
    s.takeaway('Each DoF has one owner. Other ranks keep copies where local work needs them.',
               'Ghost DoFs are copied coefficients; ghost cells are additional mesh cells (not shown here).')
    s.notes = '''Timing: 1:35.

Read the left diagram as a 2D triangular mesh, with two cell columns assigned
to rank 0 and one to rank 1. The global vertex labels run from 0 to 11. The
shared interface contains DoFs 2, 6 and 10. We deliberately assign 2 and 10 to
rank 0 and 6 to rank 1 so that forward copies travel in both directions. This
is a valid illustrative ownership convention, not a captured DOLFINx partition.
Real ownership and local ordering must come from the DoF IndexMap.

The separated diagrams show cells owned by each rank, without ghost cells.
A filled node is authoritative on this rank; an outlined node is a cached
coefficient owned remotely. Rank 0 stores eight owned entries and ghost 6;
rank 1 stores four owned entries and ghosts 2 and 10. Local storage is flat,
even though the geometric mesh is two-dimensional. Outlines retain the owner's
color, so the arrows can be followed directly from source to destination.

A forward INSERT copies current owner values into ghost slots, preserving
owned entries. Mesh ghost cells are distinct: additional cells stored locally
for algorithms requiring neighboring cell geometry/topology. DoF ghosts can be
needed by owned cells even when extra ghost cells are not shown. Reverse ADD
is a different operation that accumulates contributions into owners.

Repository: ../src/jaxghost/_metadata.py; ../src/jaxghost/jaxghost.py.
Reference: https://docs.fenicsproject.org/dolfinx/main/python/demos.html#mesh-partitioning-and-parallel-communication-analysis
'''
    out.append(s)

    s = Slide(2, 'How do we preserve this on GPUs with JAX?', 'The question',
              'Keep the FEM ownership model. Move repeated numerical work onto the devices.')
    box(s, 56, 188, 470, 226, 'MPI rank 0  /  GPU 0', color=BLUE, fill=PALE_B)
    box(s, 754, 188, 470, 226, 'MPI rank 1  /  GPU 1', color=ORANGE, fill=PALE_O)
    vector(s, 85, 270, [u(i) for i in R0], [OWNERS[i] for i in R0], ghosts=[8], cell=46, fs=19)
    vector(s, 804, 270, [u(i) for i in R1], [OWNERS[i] for i in R1], ghosts=[4,5], cell=62, fs=20)
    s.text(85, 350, 412, '9 local entries • JAX device array', 19, BLUE)
    s.text(804, 350, 385, '6 local entries • JAX device array', 19, ORANGE)
    s.arrow(540, 287, 739, 287, BLUE, 3)
    s.arrow(739, 342, 540, 342, ORANGE, 3)
    s.circle(640, 254, 30, WHITE, LINE, 1)
    s.text(611, 229, 58, '?', 53, INK, bold=True, align='center')
    box(s, 56, 451, 367, 123, 'Preserve ownership', 'Reuse DOLFINx’s owner/ghost map.', size=21)
    box(s, 456, 451, 367, 123, 'Communicate inside JIT', 'Refresh ghosts between GPU steps.', size=21)
    box(s, 857, 451, 367, 123, 'Handle unequal sizes', 'Partitions need different storage.', size=21)
    s.takeaway('How do remote owner values reach the right local ghost slots?')
    s.notes = '''Timing: 0:45.

These are exactly the local arrays from the previous slide, now held on GPUs.
We keep one MPI process per local JAX device, matching the current library's
supported layout. Rank 0 has nine local scalar entries and rank 1 has six.
Moving the vector to a GPU does not establish any ghost synchronization.

We need the same owner-to-ghost copy operation as in DOLFINx, callable within
a JIT-compiled numerical computation. DOLFINx supplies the layout at setup;
the repeated numerical values are JAX arrays. The plan must survive tracing
and account for irregular neighbor lists and unequal sizes without changing
the semantic meaning of owned and ghost entries. We will first illustrate
why matvec needs this exchange, then compare the two implemented mechanisms.

Repository: ../README.md (Current scope); ../src/jaxghost/sharded.py.
'''
    out.append(s)

    s = Slide(3, 'A local matrix row needs a remote vector value', 'Why ghosts appear',
              'A four-DoF algebraic toy • row partitioning • zero-based indices • start with y = 0')
    s.text(65, 183, 520, '1   SPLIT THE MATRIX BY OWNED ROWS', 16, MUTED, bold=True)
    # Editable 4x4 matrix with highlighted off-rank column accesses.
    vals = [[2,-1,0,0],[-1,2,-1,0],[0,-1,2,-1],[0,0,-1,2]]
    for i in range(4):
        s.rect(109, 236+i*52, 242, 49, PALE_B if i<2 else PALE_O)
        s.text(64, 247+i*52, 37, str(i), 16, BLUE if i<2 else ORANGE, align='right')
        for j in range(4):
            hi = (i,j) in [(1,2),(2,1)]
            if hi:
                s.rect(114+j*59, 239+i*52, 52, 43, WHITE, ORANGE if j==2 else BLUE, 2)
            s.text(114+j*59, 244+i*52, 52, str(vals[i][j]).replace('-', '−'), 25,
                   INK if vals[i][j] else '#A5AFBC', bold=hi, align='center')
    for x, sign in [(101,1),(357,-1)]:
        s.line(x,234,x,444,INK,2); s.line(x,234,x+sign*7,234,INK,2); s.line(x,444,x+sign*7,444,INK,2)
    s.text(125, 202, 226, 'A', 23, INK, bold=True, align='center')
    s.text(380, 315, 40, '×', 32, MUTED, align='center')
    for i,v in enumerate([1,2,4,8]):
        vector(s, 431, 239+i*52, [str(v)], [i//2], cell=51, h=43, fs=25)
    s.text(432, 202, 48, 'x', 23, INK, bold=True, align='center')
    s.text(503, 315, 34, '=', 30, MUTED)
    for i,v in enumerate([0,-1,-2,12]):
        s.text(546, 244+i*52, 65, str(v).replace('-','−'), 25, BLUE if i<2 else ORANGE, bold=True, align='center')
    s.text(546, 202, 65, 'Ax', 23, INK, bold=True, align='center')
    s.line(105, 338, 359, 338, MUTED, 1.5, dash=True)
    s.text(84, 479, 525, 'The matrix stays local; selected x values travel.', 18, MUTED)
    s.text(668, 183, 552, '2   EXCHANGE x, THEN MULTIPLY LOCAL ROWS', 16, MUTED, bold=True)
    s.rect(667, 225, 557, 155, PALE_B, radius=10)
    s.text(686, 241, 304, 'Rank 0 needs remote x₂', 19, BLUE, bold=True)
    vector(s, 1005, 241, ['1','2','4'], [0,0,1], ghosts=[2], cell=58, fs=22)
    s.text(687, 306, 509, 'y₁ = −1·1 + 2·2 − 1·4 = −1', 24, INK)
    s.rect(667, 403, 557, 155, PALE_O, radius=10)
    s.text(686, 419, 304, 'Rank 1 needs remote x₁', 19, ORANGE, bold=True)
    vector(s, 1005, 419, ['4','8','2'], [1,1,0], ghosts=[2], cell=58, fs=22)
    s.text(687, 484, 509, 'y₂ = −1·2 + 2·4 − 1·8 = −2', 24, INK)
    s.takeaway('A remote column reference becomes a local ghost-vector lookup.',
               'Our API computes owned y += Ax; x and matrix inputs are preserved, and output ghosts stay unchanged.')
    s.notes = '''Timing: 1:45.

This is a deliberately smaller algebraic toy, not the stiffness matrix of the
twelve-vertex mesh on the preceding slides. Use zero-based indices throughout.
Rank 0 owns rows and vector entries 0 and 1; rank 1 owns 2 and 3. The matrix has
2 on its diagonal and −1 on its first sub- and super-diagonals. With x equal to
[1, 2, 4, 8], the full result is [0, −1, −2, 12].

Focus on row 1: its coefficient in column 2 multiplies x₂, which lives on rank
1. Rank 0 therefore receives a copy of 4 and stores [1, 2 | 4]. Row 2 has the
symmetric need for x₁, so rank 1 stores [4, 8 | 2]. The outlined vector boxes
retain the remote owner's color. The matrix's nonlocal-column coefficients
are already stored with their owned row; it is x that needs refreshing.
An implementation remaps global columns to positions in its local extended x.

After the forward update, each rank computes its owned output rows. Setting
initial y to zero makes our owned y += Ax API look like y = Ax in this example.
The returned array preserves output ghost entries; an additional forward
exchange is needed if a later operation requires updated y ghosts. The input
x is not mutated even though an internally refreshed version feeds matvec.
Assembly's reverse accumulation is distinct from the forward exchange shown.

Repository: ../src/jaxghost/matrix.py; ../src/jaxghost/sharded_matrix.py.
'''
    out.append(s)

    s = Slide(4, 'mpi4jax: pack → sendrecv → unpack', 'First implementation',
              'The ownership plan is prepared once; the JIT-compiled operation exchanges current values.')
    box(s, 56, 189, 326, 210, '1   Pack owner entries', color=BLUE, fill=PALE_B)
    box(s, 460, 189, 360, 210, '2   Exchange with peers', color=INK)
    box(s, 898, 189, 326, 210, '3   Fill ghost slots', color=ORANGE, fill=PALE_O)
    vector(s, 81, 262, ['1','2','?'], [0,0,1], ghosts=[2], cell=56, fs=25)
    s.text(267, 270, 95, 'send 2', 20, BLUE, bold=True)
    s.text(80, 340, 278, 'x[send_indices]', 19, BLUE, font='mono')
    s.text(482, 256, 316, 'mpi4jax.sendrecv', 25, INK, font='mono', align='center')
    s.text(482, 308, 316, 'Compiled bridge → MPI', 21, MUTED, align='center')
    s.text(482, 350, 316, 'Neighbor buffers, fixed sizes', 17, MUTED, align='center')
    vector(s, 924, 262, ['1','2','4'], [0,0,1], ghosts=[2], cell=56, fs=25)
    s.text(925, 340, 274, 'ghost slots ← received', 19, ORANGE)
    s.arrow(391, 287, 450, 287, BLUE, 3)
    s.arrow(830, 287, 887, 287, ORANGE, 3)
    s.text(57, 429, 1160, 'GPU buffers can follow different transport paths', 20, INK, bold=True)
    s.rect(56, 470, 564, 102, BG, radius=8)
    s.text(74, 482, 526, 'HOST-STAGED', 14, MUTED, bold=True)
    s.text(74, 515, 526, 'GPU → CPU → MPI → CPU → GPU', 23, INK)
    s.rect(646, 470, 578, 102, BG, radius=8)
    s.text(666, 482, 536, 'CUDA-AWARE MPI', 14, MUTED, bold=True)
    s.text(666, 515, 536, 'GPU pointer → MPI → GPU pointer', 22, INK)
    s.takeaway('CUDA-aware MPI accepts device buffers; the actual transport may still stage on the host.')
    s.notes = '''Timing: 1:25.

JAXGhost builds an immutable communication layout from the DOLFINx IndexMap.
The one-time metadata phase uses mpi4py. At runtime, JAX gathers selected owned
entries into a packed buffer, loops over its sorted peers, and invokes
mpi4jax.sendrecv with the matching source, destination and buffer sizes. The
received parts are concatenated and assigned to precomputed ghost positions.
The picture follows rank 0 in the algebraic toy: send 2, receive 4.

mpi4jax lowers MPI operations into the compiled JAX computation through its
native bridge; the Python call describes an operation to execute at runtime.
Our pinned 0.9.1.post1 API returns the received array and uses ordered effects,
which also retain send-only exchanges. All ranks must participate in compatible
order with matching types and counts. Runtime communication uses a duplicated
communicator separate from host metadata traffic.

With host staging, mpi4jax copies between GPU and CPU buffers around MPI.
CUDA-aware mode instead passes device pointers to a compatible MPI stack.
That does not prove the MPI transport avoids host copies. In the recorded
IRIS investigation, both a working TCP route and a later working UCX small-
message route staged payloads inside MPI/UCX. That is a configuration-specific
observation, not a universal claim about UCX. We make no GPUDirect assertion.

Repository: ../src/jaxghost/jaxghost.py; ../HPC-GPU/transport-investigation.md.
References:
https://mpi4jax.readthedocs.io/en/latest/installation.html
https://mpi4jax.readthedocs.io/en/latest/sharp-bits.html
'''
    out.append(s)

    s = Slide(5, 'The challenge: make the whole HPC stack agree', 'What we encountered',
              'Python packages, native libraries and the selected runtime transport must be compatible.')
    # Lines first, with packages layered above them: these are compatibility edges.
    for a,b in [((223,246),(223,321)), ((490,246),(490,321)), ((990,246),(990,321)),
                ((629,211),(840,211)), ((362,355),(366,355)),
                ((863,355),(877,355)), ((223,395),(223,437)),
                ((223,437),(764,437)), ((764,437),(764,395))]:
        s.line(*a,*b, MUTED, 2)
    box(s, 56, 177, 306, 78, 'DOLFINx / PETSc', color=BLUE, fill=PALE_B, size=23)
    box(s, 366, 177, 263, 78, 'mpi4jax', color=TEAL, fill=PALE_T, size=23)
    box(s, 840, 177, 384, 78, 'JAX / jaxlib', color=ORANGE, fill=PALE_O, size=23)
    box(s, 56, 317, 306, 78, 'MPI library', color=BLUE, fill=PALE_B, size=23)
    box(s, 366, 317, 263, 78, 'mpi4py', color=TEAL, fill=PALE_T, size=23)
    box(s, 665, 317, 198, 78, 'Transport', size=21)
    box(s, 877, 317, 347, 78, 'CUDA / driver', color=ORANGE, fill=PALE_O, size=23)
    for x,w in [(70,272),(378,243),(861,342)]:
        s.rect(x,272,w,25,WHITE)
    s.text(70, 275, 272, 'shared native MPI stack', 16, MUTED, align='center')
    s.text(378, 275, 243, 'build + runtime linkage', 16, MUTED, align='center')
    s.text(861, 275, 342, 'GPU build + supported runtime', 16, MUTED, align='center')
    s.text(655, 226, 172, 'JAX compatibility', 14, MUTED, align='center')
    s.rect(330,423,323,25,WHITE)
    s.text(339, 426, 310, 'MPI selects a runtime transport', 16, MUTED, align='center')
    s.text(808, 419, 402, 'Transport must handle GPU buffers.', 16, MUTED)
    s.rect(56, 465, 564, 117, BG, radius=8)
    s.text(76, 478, 524, '01   CUDA discovery across Spack + venv', 20, INK, bold=True)
    s.text(76, 517, 524, 'mpi4py and CUDA lived in different prefixes.\nThe mpi4jax build needed a discovery fix.', 18, MUTED)
    s.rect(646, 465, 578, 117, BG, radius=8)
    s.text(666, 478, 538, '02   Selected UCX lacked GPU support', 20, INK, bold=True)
    s.text(666, 517, 538, 'The loaded transport misread a device buffer.\nA working MPI import did not validate this path.', 18, MUTED)
    s.takeaway('Successful installation is only the start: validate the GPU exchange in the actual job.')
    s.notes = '''Timing: 1:25.

The connecting lines indicate compatibility or linkage requirements, not the
order of data movement. DOLFINx and PETSc share a native MPI stack with mpi4py.
mpi4jax must use the intended MPI installation and a supported JAX/jaxlib
combination. GPU execution adds CUDA discovery, the installed GPU backend and
driver compatibility. MPI then selects a transport such as UCX at runtime.
On HPC, module order, Spack prefixes, virtual environments and launcher settings
all affect the libraries actually loaded.

Two concrete issues from our work make this less abstract. First, mpi4jax
0.9.1.post1 searched for NVIDIA wheels beside mpi4py. mpi4py was supplied by
Spack while CUDA wheels were in the virtual environment; the build missed CUDA
and needed a version-specific discovery patch. This is a historical project
example, not a claim that every current release has that bug.

Second, the selected UCX 1.15.0 installation exposed no CUDA memory domain or
CUDA modules on the GPU node. The observed send path tried a CPU access to a
device buffer and crashed. A working TCP configuration passed numerical checks
but staged internally. A subsequently tested CUDA-enabled UCX configuration
also passed; the traced small messages still used host staging. Thus the lesson
is to verify the selected stack and real exchange, not just package imports
or the top-level Open MPI CUDA-support flag. Neither issue required changing
the owner/ghost semantics.

Repository: ../HPC-GPU/notes.md (CUDA-discovery fix);
../HPC-GPU/transport-investigation.md (initial and subsequent configurations).
Reference: https://mpi4jax.readthedocs.io/en/latest/installation.html
'''
    out.append(s)

    s = Slide(6, 'Explicit ghost exchange with JAX sharding', 'Our implementation',
              'One rank per GPU • retain DOLFINx ownership • let JAX execute the numerical collective')
    s.rect(56, 169, 1168, 79, BG, radius=9)
    s.text(76, 181, 174, 'ONCE / HOST', 15, MUTED, bold=True)
    s.text(254, 179, 945, 'DOLFINx IndexMap + mpi4py metadata → packing indices, receive slots, masks', 20, INK)
    s.text(254, 211, 944, 'Initialize distributed JAX; create a one-dimensional device mesh with axis “rank”.', 17, MUTED)
    s.text(56, 271, 610, 'GLOBAL ARRAY SHAPE [2, 9]   /   ONE ROW PER GPU', 16, MUTED, bold=True)
    s.text(58, 316, 122, 'GPU 0', 21, BLUE, bold=True)
    vector(s, 191, 308, [u(i) for i in R0], [OWNERS[i] for i in R0], ghosts=[8], cell=43, fs=18)
    s.text(58, 377, 122, 'GPU 1', 21, ORANGE, bold=True)
    vector(s, 191, 369, [u(i) for i in R1]+['·']*3, [OWNERS[i] for i in R1]+[None]*3,
           ghosts=[4,5], padding=[6,7,8], cell=43, fs=18)
    s.text(190, 430, 404, 'filled = owned   |   outline = ghost   |   · = padding', 14, MUTED)
    s.rect(646, 280, 578, 170, PALE_T, radius=10)
    s.text(667, 296, 536, 'Explicit packets preserve the ghost map', 21, TEAL, bold=True)
    s.text(670, 336, 528, 'GPU 0 sends [u₂, u₁₀] → GPU 1 ghost slots', 21, BLUE)
    s.text(670, 375, 528, 'GPU 1 sends [u₆, 0*] → GPU 0 ghost slot', 21, ORANGE)
    s.text(670, 412, 528, '* Fixed-width packet padding is masked / discarded.', 16, MUTED)
    s.text(57, 474, 1150, 'EACH JIT-COMPILED CALL   /   shard_map runs the same body on every device', 16, MUTED, bold=True)
    blocks = [(56, 230, 'Pack by indices'), (320, 314, 'lax.all_to_all'), (668, 264, 'Scatter to ghosts'), (966, 258, 'Local CSR matvec')]
    for x,w,t in blocks:
        s.rect(x, 514, w, 62, PALE_T if 'all_to_all' in t else BG, TEAL if 'all_to_all' in t else LINE, 1.5, radius=8)
        s.text(x+12, 531, w-24, t, 21, TEAL if 'all_to_all' in t else INK,
               font='mono' if 'all_to_all' in t else 'bold', align='center')
    for x in [292,640,938]:
        s.arrow(x,545,x+21,545,TEAL,2)
    s.takeaway('Sharding places the arrays. Our plan defines which owner values fill which ghosts.',
               'Current scope: forward INSERT + scalar CSR matvec. Tradeoff: padding and a global all-to-all collective.')
    s.notes = '''Timing: 2:25. Elapsed speaking time through this slide: 10:00.

Reuse the twelve-vertex mesh's exact local order. Rank 0 has eight owned values
and one ghost: nine entries. Rank 1 has four owned values and two ghosts: six
entries, padded to nine. The conceptual global array has shape [2, 9], with its
first axis sharded across the two devices. It is not a dense representation of
the original global FEM vector: the rows contain local storage including ghost
duplicates and padding. Each process holds its own addressable row, not both.

The setup phase uses DOLFINx IndexMap ownership and mpi4py metadata exchanges.
It computes globally consistent padded lengths and fixed-width peer packets,
then places send indices, receive positions and validity masks into sharded
JAX arrays. Initialize distributed JAX before querying devices and align mesh
positions with MPI ranks. DOLFINx is used at preparation, and is not needed for
replay of already exported fixtures; mpi4py remains used for setup/bootstrap.

At runtime, shard_map runs the same device body on each shard, with rank-specific
tables as operands. Gather the required owner values and mask unused entries.
The native lax.all_to_all collective routes fixed-width packets by destination.
Rank 0 sends u₂ and u₁₀; rank 1 sends u₆ plus one padded item. Self and unused
peer slots also exist in the all-to-all layout. Receive-position tables route
valid data into the original ghost slots. Invalid positions use an out-of-bounds
sentinel with scatter mode='drop'. Owned values and array padding are preserved.

ShardedJAXMatrixCSR.mult then computes owned y += Ax using local scalar CSR
storage. It preserves the caller's arrays and output ghosts, and does not
overlap communication with multiplication. There are two distinct paddings:
the vector padding drawn in gray, and fixed-width peer packet padding. This
implementation uses all_to_all, not ragged_all_to_all or an automatic halo
inference mechanism. The padded representation and all-rank collective cost
are the tradeoff for fixed shapes and native JAX communication.

The repeated numerical exchange no longer uses the mpi4jax bridge or a runtime
MPI communicator in our operator. This does not remove all MPI dependencies
from setup or establish a physical GPU transport. Forward INSERT and scalar
CSR matvec are implemented; reverse ADD is provided by the MPI backend but not
this sharded backend. Differentiation is outside the supported interface.

Repository: ../src/jaxghost/sharded.py; ../src/jaxghost/sharded_matrix.py;
../datasets/test-sharded-gpu/results/matrix-timing-20260909T052640Z/REPORT.md.
Reference: https://docs.jax.dev/en/latest/notebooks/shard_map.html
'''
    out.append(s)
    s = Slide(7, 'Newer JAX APIs: optimization opportunities', 'What to try next',
              'Candidates for our sharded implementation • preserve the ownership map • benchmark on the target GPUs')
    for x, title, col, fill in [(56, 'Less packet padding', BLUE, PALE_B),
                                (456, 'Reuse device storage', TEAL, PALE_T),
                                (857, 'Overlap useful work', ORANGE, PALE_O)]:
        s.rect(x, 181, 367, 346, fill, LINE, 1, radius=10)
        s.text(x+20, 200, 327, title, 23, col, bold=True)
    s.text(76, 249, 327, 'lax.ragged_all_to_all', 20, BLUE, font='mono')
    s.text(76, 294, 327, 'Fixed peer packets', 16, MUTED)
    for i in range(8):
        s.rect(76+i*34, 321, 28, 22, BLUE if i in [0,4,5] else WHITE,
               BLUE if i in [0,4,5] else LINE, 1)
    s.arrow(365, 330, 365, 375, BLUE, 2, head=8)
    s.text(76, 361, 280, 'Packed values + counts / offsets', 16, MUTED)
    for i in range(3):
        s.rect(76+i*34, 389, 28, 22, BLUE)
    s.text(76, 440, 327, 'Send variable-length slices.\nKeep fixed buffer capacities;\nvector / CSR padding remains.', 18, INK)

    s.text(476, 249, 327, 'donate_argnums / jax.new_ref', 18, TEAL, font='mono')
    vector(s, 491, 316, ['1','2','4'], [0,0,1], ghosts=[2], cell=62, fs=22)
    s.text(688, 325, 112, '← write', 18, ORANGE, bold=True)
    s.text(476, 382, 327, 'Reuse buffers or write ghost slots.', 17, MUTED)
    s.text(476, 440, 327, 'Donation consumes the input.\nRefs enable explicit mutation.\nMeasure memory + dispatch cost.', 18, INK)

    s.text(877, 249, 327, 'lax.psend / lax.precv', 20, ORANGE, font='mono')
    s.text(877, 294, 327, 'Target schedule — verify in a trace', 16, MUTED)
    s.rect(877, 327, 235, 28, ORANGE, radius=4)
    s.text(885, 331, 216, 'communicate ghost values', 14, WHITE, bold=True, align='center')
    s.rect(877, 370, 192, 28, BLUE, radius=4)
    s.text(884, 374, 177, 'owned-column part', 14, WHITE, bold=True, align='center')
    s.rect(1119, 370, 84, 28, ORANGE, radius=4)
    s.text(1123, 374, 76, 'ghosts', 14, WHITE, bold=True, align='center')
    s.arrow(1112, 341, 1126, 366, ORANGE, 1.5, head=6)
    s.text(877, 440, 327, 'Split owned / ghost-column work.\nTest scheduling and GPU progress.\nMPI-style overlap is not automatic.', 18, INK)
    s.text(57, 548, 1165, 'Also profile local CSR alternatives; consider Pallas fusion only for a measured kernel bottleneck.', 19, MUTED)
    s.takeaway('First test storage reuse and ragged exchange; measure end-to-end matvec before choosing.',
               'Optimization candidates, not measured speedups. Verify GPU execution, numerical results and required JAX transforms.')
    s.notes = '''Timing: 1:30. Total planned speaking time with the extra slide: 11:30.

This is a roadmap, not a report of optimizations already implemented. These
are capabilities in current JAX documentation; they were not all introduced
in the same release. The supplied audit reports small CPU probes on JAX 0.10.2,
not complete GPU or transformation validation. We did not rerun those probes
or change the numerical backend while preparing this slide. Check the actual
JAX/jaxlib and GPU stack before each experiment.

LESS PADDING. Our existing all-to-all allocates one maximum-width numerical
packet per peer. The mini-diagram illustrates the same three useful values
inside eight packet slots versus a packed representation; it is not a measured
size reduction for the earlier two-rank mesh. ragged_all_to_all exchanges slices
using counts and offsets, matching the DOLFINx-derived plan. Zero counts can
represent non-neighbors. With our uniform shards, capacities remain fixed:
use the maximum total outgoing count rather than rank count times the largest
peer message. Metadata and vector/CSR padding remain. output_offsets specifies
where each outgoing slice lands on its receiver, so receiver offsets must be
exchanged during setup before unpacking into original ghost order. Validate
GPU lowering and compare latency, not just payload size; the supplied audit's
CPU probe could not compile this primitive.

REUSE STORAGE. Donation is the smallest API experiment: permit reuse of the
forward input x or matvec accumulator y when the caller relinquishes it. The
donated input must not be reused afterward; memory reuse is an opportunity
for the compiler, not a promised speedup. Keep the existing preserving API as
the baseline. Mutable Refs offer a separate API with indexed writes to ghost
or owned-output slots. Pass Refs explicitly into shard_map, rather than closing
over them, and perform the indexed updates directly. Wrapping the existing
full-vector functional update in a Ref will not automatically remove its
temporaries. Current Ref documentation notes slower Python dispatch to impure
JIT functions taking Ref inputs; benchmark against donated functional calls.

OVERLAP. Split the local operator into owned-column and ghost-column parts.
The timeline is a target: initiate the exchange, compute the owned-column part,
then use received values for the ghost-column part. psend/precv expose separate
send/receive operations; their GPU semantics map to NCCL communication. They
are not interchangeable with MPI nonblocking begin/end. Matching permutations,
fixed operand shapes, ordering and progress need testing. The supplied audit
found no differentiation rule for psend; do not infer AD support from JIT
support. Also compare ordinary collectives with independent computation and
XLA latency-hiding / profile-guided scheduling. A GPU trace must show overlap.

FOLLOW-UP. Profile local CSR separately: the current JAX sparse module is
experimental and does not promise performance-critical suitability. Consider
BCSR for batched vectors or gather-plus-segment-sum with cached row indices,
keeping existing CSR as the baseline. Pallas/Mosaic GPU offers later kernel-
level communication and computation integration. Its dense collective matmul
example does not establish performance or hardware suitability for our
irregular float64 CSR workload, including the GPUs used in this project.

Suggested order: global no-communication fast path and donation; persistent
Refs; GPU ragged exchange; profile local kernels; test overlap; specialized
kernels only if warranted. The no-communication decision must be global:
a rank with no local ghosts may still need to send to another rank. Retain the
DOLFINx ownership and ghost ordering, and check correctness, memory and total
matvec time for each candidate. Backend execution, JIT, batching and AD are
separate properties to verify.

References (official JAX documentation, checked for this slide):
https://docs.jax.dev/en/latest/_autosummary/jax.lax.ragged_all_to_all.html
https://docs.jax.dev/en/latest/buffer_donation.html
https://docs.jax.dev/en/latest/array_refs.html
https://docs.jax.dev/en/latest/_autosummary/jax.lax.psend.html
https://docs.jax.dev/en/latest/_autosummary/jax.lax.precv.html
https://docs.jax.dev/en/latest/gpu_performance_tips.html
https://docs.jax.dev/en/latest/jax.experimental.sparse.html
https://docs.jax.dev/en/latest/pallas/gpu/collective_matmul.html

Repository baseline: ../src/jaxghost/sharded.py;
../src/jaxghost/sharded_matrix.py. The user-supplied API audit informed the
priorities; its probe outcomes are attributed above rather than revalidated.
'''
    out.append(s)
    return out


def color(s):
    return RGBColor.from_string(s.lstrip('#'))


def pptx_render(deck, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(W*SCALE), Pt(H*SCALE)
    prs.core_properties.title = 'From distributed FEM to JAX GPU ghost exchange'
    prs.core_properties.subject = 'DOLFINx ownership, mpi4jax and explicit JAX sharded halos'
    prs.core_properties.author = 'JAX-ghost project'
    def unit(v): return Pt(v*SCALE)
    for d in deck:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for o in d.objects:
            k = o['kind']
            if k == 'text':
                shape = slide.shapes.add_textbox(unit(o['x']),unit(o['y']),unit(o['w']),unit(o['h']))
                tf = shape.text_frame
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                tf.word_wrap=False
                tf.vertical_anchor=MSO_ANCHOR.TOP
                for i, line in enumerate(o['text'].split('\n')):
                    p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
                    p.text = line
                    p.alignment = {'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[o['align']]
                    p.space_before=p.space_after=Pt(0)
                    p.line_spacing = 1.3
                    p.font.name = 'DejaVu Sans Mono' if o['font']=='mono' else 'DejaVu Sans'
                    p.font.size=unit(o['size'])
                    p.font.bold=o['font']=='bold'
                    p.font.color.rgb=color(o['color'])
            elif k == 'line':
                # Freeform handles any direction without flips or connector heuristics.
                b=slide.shapes.build_freeform(o['x1'],o['y1'],scale=unit(1))
                b.add_line_segments([(o['x2'],o['y2'])],close=False)
                shape=b.convert_to_shape()
                shape.fill.background()
                shape.line.color.rgb=color(o['color']); shape.line.width=unit(o['lw'])
                if o.get('dash'):
                    from pptx.enum.dml import MSO_LINE_DASH_STYLE
                    shape.line.dash_style=MSO_LINE_DASH_STYLE.DASH
            else:
                if k=='poly':
                    b=slide.shapes.build_freeform(*o['points'][0],scale=unit(1))
                    b.add_line_segments(o['points'][1:],close=True)
                    shape=b.convert_to_shape()
                else:
                    typ=MSO_SHAPE.OVAL if k=='ellipse' else MSO_SHAPE.ROUNDED_RECTANGLE if o.get('radius') else MSO_SHAPE.RECTANGLE
                    shape=slide.shapes.add_shape(typ,unit(o['x']),unit(o['y']),unit(o['w']),unit(o['h']))
                    if o.get('radius'):
                        shape.adjustments[0]=o['radius']/min(o['w'],o['h'])
                if o.get('fill'):
                    shape.fill.solid(); shape.fill.fore_color.rgb=color(o['fill'])
                else: shape.fill.background()
                if o.get('stroke'):
                    shape.line.color.rgb=color(o['stroke']); shape.line.width=unit(o['lw'])
                else: shape.line.fill.background()
            shape.name = f'{d.number:02d}-{len(slide.shapes):03d}-{k}'
        slide.notes_slide.notes_text_frame.text = d.title+'\n\n'+d.notes.strip()
    prs.save(path)


def pdf_render(deck, path):
    c=canvas.Canvas(str(path),pagesize=(W*SCALE,H*SCALE),pageCompression=1)
    c.setTitle('From distributed FEM to JAX GPU ghost exchange')
    c.setAuthor('JAX-ghost project')
    for d in deck:
        c.saveState(); c.scale(SCALE,SCALE)
        for o in d.objects:
            k=o['kind']; c.saveState()
            if k=='text':
                c.setFillColor(HexColor(o['color'])); c.setFont(o['font'],o['size'])
                ascent=pdfmetrics.getAscent(o['font'])*o['size']/1000
                for i,line in enumerate(o['text'].split('\n')):
                    yy=H-o['y']-ascent-i*o['size']*1.3
                    if o['align']=='center': c.drawCentredString(o['x']+o['w']/2,yy,line)
                    elif o['align']=='right': c.drawRightString(o['x']+o['w'],yy,line)
                    else: c.drawString(o['x'],yy,line)
            elif k=='line':
                c.setStrokeColor(HexColor(o['color'])); c.setLineWidth(o['lw'])
                if o.get('dash'): c.setDash(5,4)
                c.line(o['x1'],H-o['y1'],o['x2'],H-o['y2'])
            else:
                if o.get('fill'): c.setFillColor(HexColor(o['fill']))
                if o.get('stroke'): c.setStrokeColor(HexColor(o['stroke']))
                c.setLineWidth(o['lw']); opts={'stroke':int(bool(o.get('stroke'))),'fill':int(bool(o.get('fill')))}
                if k=='poly':
                    p=c.beginPath(); p.moveTo(o['points'][0][0],H-o['points'][0][1])
                    for x,y in o['points'][1:]: p.lineTo(x,H-y)
                    p.close(); c.drawPath(p,**opts)
                elif k=='ellipse': c.ellipse(o['x'],H-o['y']-o['h'],o['x']+o['w'],H-o['y'],**opts)
                elif o.get('radius'): c.roundRect(o['x'],H-o['y']-o['h'],o['w'],o['h'],o['radius'],**opts)
                else: c.rect(o['x'],H-o['y']-o['h'],o['w'],o['h'],**opts)
            c.restoreState()
        c.restoreState(); c.showPage()
    c.save()


def validate_and_preview(deck, pptx_path, pdf_path):
    # Validate the exact illustrative mesh and toy independently of JAX runtimes.
    own0, own1 = set(R0[:8]), set(R1[:4])
    assert own0.isdisjoint(own1) and own0|own1==set(range(12))
    assert R0[8:]==[6] and R1[4:]==[2,10]
    assert all(OWNERS[i]==0 for i in own0) and all(OWNERS[i]==1 for i in own1)
    assert {i for i in R0 if OWNERS[i]==1}=={6}
    assert {i for i in R1 if OWNERS[i]==0}=={2,10}
    A=[[2,-1,0,0],[-1,2,-1,0],[0,-1,2,-1],[0,0,-1,2]]
    x=[1,2,4,8]
    y=[sum(a*b for a,b in zip(row,x)) for row in A]
    assert y==[0,-1,-2,12]
    local0=[x[0],x[1],x[2]]; local1=[x[2],x[3],x[1]]
    assert -local0[0]+2*local0[1]-local0[2]==-1
    assert -local1[2]+2*local1[0]-local1[1]==-2
    prs=Presentation(pptx_path); doc=fitz.open(pdf_path)
    assert len(prs.slides)==len(doc)==len(deck)==SLIDE_COUNT
    for i,(d,ps,page) in enumerate(zip(deck,prs.slides,doc)):
        assert len(ps.shapes)==len(d.objects)
        expected_text=[o['text'] for o in d.objects if o['kind']=='text']
        pptx_text=[sh.text for sh in ps.shapes if sh.name.endswith('-text')]
        assert pptx_text==expected_text, 'PowerPoint text must match shared definitions'
        pdf_text=' '.join(page.get_text().split())
        for t in expected_text:
            assert ' '.join(t.split()) in pdf_text, f'Missing PDF text: {t}'
        assert d.title in ps.notes_slide.notes_text_frame.text
        assert not page.get_images(), 'PDF must remain vector-based'
        assert len(ps.shapes)>20
        for o in d.objects:
            if o['kind']=='text':
                assert 0<=o['x'] and o['x']+o['w']<=W+1
                assert 0<=o['y'] and o['y']+o['h']<=H+1
        for b in page.get_text('dict')['blocks']:
            for line in b.get('lines',[]):
                for sp in line['spans']:
                    xx,yy,xx2,yy2=sp['bbox']
                    assert xx>=0 and yy>=0 and xx2<=W*SCALE+1 and yy2<=H*SCALE+1, sp
        assert str(i) in page.get_text()
    previews=ROOT/'previews'; previews.mkdir(exist_ok=True)
    contact=Image.new('RGB',(1280,math.ceil(len(deck)/2)*385),'#E5EAF0')
    draw=ImageDraw.Draw(contact)
    for i,p in enumerate(doc):
        pix=p.get_pixmap(matrix=fitz.Matrix(1.6,1.6),alpha=False)
        pix.save(previews/f'slide-{i}.png')
        thumb=Image.open(previews/f'slide-{i}.png').resize((608,342),Image.Resampling.LANCZOS)
        xx,yy=16+(i%2)*640,16+(i//2)*385
        contact.paste(thumb,(xx,yy))
        draw.text((xx,yy+348),f'SLIDE {i}  /  {deck[i].title}',fill=INK,
                  font=ImageFont.truetype(str(FONTS/'DejaVuSans.ttf'),11))
    contact.save(ROOT/'overview.png')
    return {'slides':len(deck),'pdf_pages':len(doc),'editable_shapes':sum(len(s.shapes) for s in prs.slides),
            'raster_images_in_pdf':0,'toy_result':y,'mesh_owned_counts':[8,4],
            'mesh_ghost_counts':[1,2],'padded_shard_shape':[2,9],
            'text_geometry':'all text boxes and PDF spans within canvas',
            'rendering':'PDF rendered from shared geometry; PowerPoint structure checked',
            'runtime_semantics':'reviewed against existing source; no numerical backend modified'}


def main():
    deck=slides()
    pptx_path=ROOT/'jax-ghost-presentation.pptx'
    pdf_path=ROOT/'jax-ghost-presentation.pdf'
    pptx_render(deck,pptx_path); pdf_render(deck,pdf_path)
    (ROOT/'speaker-notes.md').write_text('# From distributed FEM to JAX GPU ghost exchange\n\n'
        'Eight slides, numbered 0–7. Target duration: 11½ minutes.\n\n'
        + '\n\n'.join(f'## Slide {s.number} — {s.title}\n\n{s.notes.strip()}' for s in deck)+'\n')
    report=validate_and_preview(deck,pptx_path,pdf_path)
    (ROOT/'validation.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))


if __name__=='__main__':
    main()
